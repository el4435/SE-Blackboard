"""Convert Response_to_Reviewers.md to a Word document.

Handles: headers, paragraphs, bold (**...**), inline code (`...`), simple
markdown tables, and bullet lists. Strips LaTeX math markers like $...$
to plain text inside the conversion (Word formula rendering is out of
scope for this revision).
"""

from __future__ import annotations

import os
import re

from docx import Document
from docx.shared import Pt

SRC = r"E:\SE-Blackboard\paper\Response_to_Reviewers.md"
DST = r"E:\SE-Blackboard\paper\Response_to_Reviewers.docx"


def add_runs(p, text: str):
    # Inline parsing: **bold**, `code`, $math$ (strip $)
    text = re.sub(r"\$([^$]*)\$", r"\1", text)
    pieces = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\\cite\{[^}]*\})", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = p.add_run(piece[2:-2])
            r.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = p.add_run(piece[1:-1])
            r.font.name = "Consolas"
        elif piece.startswith("\\cite{"):
            r = p.add_run("[" + piece[6:-1] + "]")
        else:
            p.add_run(piece)


def main() -> None:
    with open(SRC, encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")
        if not line.strip():
            i += 1
            continue

        # Header
        m = re.match(r"^(#+)\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2), level=level)
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ("---", "***"):
            doc.add_paragraph()
            i += 1
            continue

        # Table: a line of pipes followed by a separator line of dashes
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[-:|\s]+\|$", lines[i + 1].rstrip()):
            header_cells = [c.strip() for c in line.strip("|").split("|")]
            i += 2  # skip header + separator
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row_cells = [c.strip() for c in lines[i].rstrip("\n").strip("|").split("|")]
                rows.append(row_cells)
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(header_cells))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(header_cells):
                cell = t.rows[0].cells[j]
                cell.text = ""
                add_runs(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for ri, row in enumerate(rows, 1):
                for j, c in enumerate(row[: len(header_cells)]):
                    cell = t.rows[ri].cells[j]
                    cell.text = ""
                    add_runs(cell.paragraphs[0], c)
            doc.add_paragraph()
            continue

        # Bullet list
        if re.match(r"^\s*-\s+", line):
            while i < len(lines) and re.match(r"^\s*-\s+", lines[i]):
                content = re.sub(r"^\s*-\s+", "", lines[i].rstrip("\n"))
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, content)
                i += 1
            continue

        # Numbered list
        if re.match(r"^\s*\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                content = re.sub(r"^\s*\d+\.\s+", "", lines[i].rstrip("\n"))
                p = doc.add_paragraph(style="List Number")
                add_runs(p, content)
                i += 1
            continue

        # Plain paragraph (may span multiple consecutive non-empty lines)
        parts = [line]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not re.match(r"^(#+\s|---|\||\s*-\s|\s*\d+\.\s)", lines[i])
        ):
            parts.append(lines[i].rstrip("\n"))
            i += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(parts))

    doc.save(DST)
    print(f"Wrote {DST} ({os.path.getsize(DST):,} bytes)")


if __name__ == "__main__":
    main()
