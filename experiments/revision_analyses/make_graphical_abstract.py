"""Build a Graphical Abstract for IEEE Access final-files submission.

Source: paper/ACCESS_latex_template_20240429/figures/fig4_bottleneck_model.png
Target: paper/GA.jpg, 660x295, JPEG, <= 45 KB

Strategy: open the source PNG, fit-and-letterbox to a 660x295 canvas
(white background), save as JPEG with progressively lower quality until
the file size is under 45 KB. Also writes paper/GA_caption.docx with the
filename reference and the <=60-word caption.
"""

from __future__ import annotations

import io
import os

from PIL import Image
from docx import Document


ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

SRC = os.path.join(
    ROOT, "paper", "ACCESS_latex_template_20240429", "figures",
    "fig4_bottleneck_model.png",
)
DST_JPG = os.path.join(ROOT, "paper", "GA.jpg")
DST_DOC = os.path.join(ROOT, "paper", "GA_caption.docx")

TARGET_W, TARGET_H = 660, 295
SIZE_LIMIT_BYTES = 45 * 1024

CAPTION = (
    "In multi-agent software-engineering systems built on large language "
    "models, the way agents share information strongly affects pipeline "
    "quality, yet this design choice is rarely studied. A shared-state "
    "architecture preserves far more technical detail than conventional "
    "message passing and improves bug-localization accuracy, but a downstream "
    "code-editing bottleneck still limits how often software issues are "
    "resolved."
)


def fit_with_padding(src: Image.Image, w: int, h: int) -> Image.Image:
    src = src.convert("RGB")
    sw, sh = src.size
    ratio = min(w / sw, h / sh)
    nw, nh = int(round(sw * ratio)), int(round(sh * ratio))
    resized = src.resize((nw, nh), Image.LANCZOS)
    canvas = Image.new("RGB", (w, h), color=(255, 255, 255))
    canvas.paste(resized, ((w - nw) // 2, (h - nh) // 2))
    return canvas


def save_under_limit(img: Image.Image, path: str, limit: int) -> int:
    for q in (90, 85, 80, 75, 70, 65, 60, 55, 50, 45, 40, 35):
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=q, optimize=True, progressive=True)
        size = buf.tell()
        if size <= limit:
            with open(path, "wb") as f:
                f.write(buf.getvalue())
            return q
    # Fallback: write at lowest quality even if over limit
    img.save(path, format="JPEG", quality=35, optimize=True, progressive=True)
    return 35


def main() -> None:
    src = Image.open(SRC)
    print(f"Source: {src.size} {src.mode}")
    target = fit_with_padding(src, TARGET_W, TARGET_H)
    q = save_under_limit(target, DST_JPG, SIZE_LIMIT_BYTES)
    final_size = os.path.getsize(DST_JPG)
    print(f"Wrote {DST_JPG} -> {final_size:,} bytes at quality {q}")

    # Caption word count check
    n_words = len(CAPTION.split())
    print(f"Caption word count: {n_words}")
    if n_words > 60:
        print("WARNING: caption exceeds 60 words!")

    # Build GA caption Word doc
    doc = Document()
    doc.add_heading("Graphical Abstract Submission", level=1)
    doc.add_paragraph(f"GA filename: GA.jpg ({TARGET_W}x{TARGET_H}, {final_size:,} bytes)")
    doc.add_paragraph(f"Overlay/still filename: (not applicable; no video submission)")
    doc.add_heading("Caption", level=2)
    doc.add_paragraph(CAPTION)
    doc.save(DST_DOC)
    print(f"Wrote {DST_DOC}")


if __name__ == "__main__":
    main()
